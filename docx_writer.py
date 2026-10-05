# Module for appending text or HTML into project .docx files
import json
import os
import shutil
import subprocess
import tempfile
import threading
import time
from datetime import datetime

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from config import PENDING_DIR
from config import load_config
from logger import log
from platform_utils import is_pandoc_available, get_pandoc_path


# Hide console window when spawning pandoc on Windows
CREATE_NO_WINDOW = 0x08000000 if os.name == "nt" else 0


# Writer class for project .docx files
class DocxWriter:
    # Initialize writer
    def __init__(self):
        self._lock_docx = False
        self.last_appended_paragraphs = 0
        self._doc = None
        self._doc_path = None
        self._dirty = False
        self._pending_paragraphs = 0
        self._flush_timer = None
        self._flush_lock = threading.Lock()

    # Return cached Document for path, opening on first use
    def _get_doc(self, docx_path: str):
        if self._doc is not None and self._doc_path == docx_path:
            return self._doc
        if self._doc is not None and self._dirty:
            self._flush_now()
        if os.path.exists(docx_path):
            self._doc = Document(docx_path)
        else:
            self._doc = Document()
        self._doc_path = docx_path
        return self._doc

    # Mark Document as dirty and schedule a delayed flush
    def _mark_dirty(self, added: int) -> None:
        self._dirty = True
        self._pending_paragraphs += max(0, int(added))
        if self._flush_timer is not None:
            self._flush_timer.cancel()
        self._flush_timer = threading.Timer(5.0, self._flush_now)
        self._flush_timer.daemon = True
        self._flush_timer.start()

    # Write buffered Document to disk
    def _flush_now(self) -> None:
        with self._flush_lock:
            if not self._dirty or self._doc is None or not self._doc_path:
                return
            try:
                self._doc.save(self._doc_path)
                log.info("DocxWriter: flushed %d paragraphs into %s", self._pending_paragraphs, self._doc_path)
                self._dirty = False
                self._pending_paragraphs = 0
            except PermissionError as e:
                log.warning("DocxWriter: flush busy %s: %s", self._doc_path, e)
            except Exception as e:
                log.exception("DocxWriter: flush error %s: %s", self._doc_path, e)

    # Public flush for exit / project change
    def flush(self) -> None:
        self._flush_now()

    # Close writer: cancel timer, flush, drop cached Document
    def close(self) -> None:
        if self._flush_timer is not None:
            self._flush_timer.cancel()
            self._flush_timer = None
        self._flush_now()
        self._doc = None
        self._doc_path = None

    # Append plain text paragraph to .docx
    def append_text(self, docx_path: str, text: str) -> bool:
        if not docx_path:
            log.error("DocxWriter: empty .docx path")
            return False
        folder = os.path.dirname(docx_path)
        if folder and not os.path.isdir(folder):
            try:
                os.makedirs(folder, exist_ok=True)
            except Exception as e:
                log.exception("DocxWriter: cannot create folder %s: %s", folder, e)
                return False
        try:
            doc = self._get_doc(docx_path)
            before = len(doc.paragraphs)
            self._add_stamp(doc)
            doc.add_paragraph(text)
            doc.add_paragraph("")
            self.last_appended_paragraphs = len(doc.paragraphs) - before
            self._mark_dirty(len(doc.paragraphs) - before)
            log.info("DocxWriter: buffered %d chars for %s", len(text), docx_path)
            return True
        except PermissionError as e:
            log.warning("DocxWriter: file busy %s: %s", docx_path, e)
            self._enqueue(docx_path, text, item_type="text")
            return False
        except Exception as e:
            log.exception("DocxWriter: write error in %s: %s", docx_path, e)
            return False

    # Append HTML fragment into .docx preserving structure
    def append_html(self, docx_path: str, html_text: str) -> bool:
        if not docx_path or not html_text:
            log.error("DocxWriter: empty .docx path or empty HTML")
            return False
        folder = os.path.dirname(docx_path)
        if folder and not os.path.isdir(folder):
            try:
                os.makedirs(folder, exist_ok=True)
            except Exception as e:
                log.exception("DocxWriter: cannot create folder %s: %s", folder, e)
                return False
        from html_to_docx import convert
        parsed = convert(html_text)
        if parsed is None:
            log.warning("DocxWriter: HTML conversion failed, skipping")
            return False
        try:
            doc = self._get_doc(docx_path)
            before = len(doc.paragraphs)
            self._add_stamp(doc)
            self._copy_paragraphs(parsed, doc)
            doc.add_paragraph("")
            self._mark_dirty(len(doc.paragraphs) - before)
            log.info("DocxWriter: buffered HTML fragment for %s", docx_path)
            return True
        except PermissionError as e:
            log.warning("DocxWriter: file busy %s: %s", docx_path, e)
            self._enqueue(docx_path, html_text, item_type="html")
            return False
        except Exception as e:
            log.exception("DocxWriter: HTML write error in %s: %s", docx_path, e)
            return False

    # Append Markdown fragment into .docx preserving structure
    def append_markdown(self, docx_path: str, markdown_text: str) -> bool:
        if not docx_path or not markdown_text:
            log.error("DocxWriter: empty .docx path or empty markdown")
            return False
        folder = os.path.dirname(docx_path)
        if folder and not os.path.isdir(folder):
            try:
                os.makedirs(folder, exist_ok=True)
            except Exception as e:
                log.exception("DocxWriter: cannot create folder %s: %s", folder, e)
                return False
        from markdown_to_docx import append_markdown as md_append
        try:
            doc = self._get_doc(docx_path)
            before = len(doc.paragraphs)
            self._add_stamp(doc)
            if self._pandoc_enabled():
                count, chars_written = self._append_markdown_pandoc(doc, markdown_text)
                if count < 0:
                    log.info("DocxWriter: pandoc failed, fallback to mistune")
                    count, chars_written = md_append(doc, markdown_text)
            else:
                count, chars_written = md_append(doc, markdown_text)
            if count <= 0:
                log.warning("DocxWriter: markdown produced no blocks, fallback to plain text")
                doc.add_paragraph(markdown_text)
                chars_written = len(markdown_text)
            self._check_loss(len(markdown_text), chars_written)
            doc.add_paragraph("")
            self.last_appended_paragraphs = len(doc.paragraphs) - before
            self._mark_dirty(self.last_appended_paragraphs)
            log.info("DocxWriter: buffered markdown (%d blocks, %d chars, %d paragraphs) for %s", count, chars_written, self.last_appended_paragraphs, docx_path)
            return True
        except PermissionError as e:
            log.warning("DocxWriter: file busy %s: %s", docx_path, e)
            self._enqueue(docx_path, markdown_text, item_type="markdown")
            return False
        except Exception as e:
            log.exception("DocxWriter: markdown write error in %s: %s", docx_path, e)
            return False

    # Compare source and rendered char counts, log on significant loss
    def _check_loss(self, source_len: int, written_len: int) -> None:
        if source_len <= 0:
            return
        threshold = 0.10
        min_abs = 20
        try:
            from config import load_config
            cfg = load_config()
            threshold = float(getattr(cfg, "loss_threshold_pct", 0.10))
            min_abs = int(getattr(cfg, "loss_min_abs_chars", 20))
        except Exception:
            pass
        lost = source_len - written_len
        pct = lost / source_len
        # Both percent and absolute loss must exceed their limits to warn/error
        if pct >= threshold and lost >= min_abs:
            if pct >= 2 * threshold:
                log.error("DocxWriter: high content loss %d of %d chars (%.1f%%) for source", lost, source_len, pct * 100)
            else:
                log.warning("DocxWriter: content loss %d of %d chars (%.1f%%) for source", lost, source_len, pct * 100)
        else:
            log.debug("DocxWriter: content delta %d of %d chars (%.2f%%)", lost, source_len, pct * 100)

    # Check whether pandoc is enabled in config and available in PATH
    def _pandoc_enabled(self) -> bool:
        try:
            cfg = load_config()
        except Exception:
            return False
        if not getattr(cfg, "use_pandoc", False):
            return False
        return is_pandoc_available()

    # Convert markdown to docx via pandoc, then copy paragraphs into target doc
    # Returns (paragraphs_added, chars_written) or (-1, 0) on failure
    def _append_markdown_pandoc(self, doc, markdown_text: str) -> tuple:
        pandoc = get_pandoc_path()
        if not pandoc:
            return -1, 0
        tmp_dir = tempfile.mkdtemp(prefix="chatclipper_pandoc_")
        md_path = os.path.join(tmp_dir, "in.md")
        out_path = os.path.join(tmp_dir, "out.docx")
        try:
            with open(md_path, "w", encoding="utf-8") as f:
                f.write(markdown_text)
            proc = subprocess.run(
                [pandoc, md_path, "-o", out_path, "--from=markdown", "--to=docx"],
                capture_output=True, timeout=30, creationflags=CREATE_NO_WINDOW,
            )
            if proc.returncode != 0:
                log.warning("DocxWriter: pandoc failed rc=%d", proc.returncode)
                return -1, 0
            tmp_doc = Document(out_path)
            before = len(doc.paragraphs)
            self._copy_paragraphs(tmp_doc, doc)
            added = len(doc.paragraphs) - before
            chars = sum(len(p.text) for p in tmp_doc.paragraphs)
            return added, chars
        except Exception as e:
            log.warning("DocxWriter: pandoc error: %s", e)
            return -1, 0
        finally:
            try:
                shutil.rmtree(tmp_dir, ignore_errors=True)
            except Exception:
                pass

    # Copy all paragraphs from source document into target document
    def _copy_paragraphs(self, source, target) -> None:
        for p in source.paragraphs:
            new_p = target.add_paragraph()
            try:
                if p.style is not None:
                    new_p.style = p.style
            except Exception:
                pass
            for r in p.runs:
                new_run = new_p.add_run(r.text)
                try:
                    new_run.bold = r.bold
                    new_run.italic = r.italic
                    new_run.underline = r.underline
                except Exception:
                    pass

    # Add timestamp stamp with horizontal line above it
    # Format: dd.mm.yy hh:mm
    def _add_stamp(self, doc) -> None:
        try:
            self._add_horizontal_line(doc)
            stamp = datetime.now().strftime("%d.%m.%y %H:%M")
            doc.add_paragraph(stamp)
        except Exception as e:
            log.exception("DocxWriter: stamp error: %s", e)

    # Add an empty paragraph with bottom border as a horizontal line
    def _add_horizontal_line(self, doc) -> None:
        try:
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            pPr.append(pBdr)
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "auto")
            pBdr.append(bottom)
        except Exception as e:
            log.exception("DocxWriter: horizontal line error: %s", e)

    # Queue file path for a specific .docx
    def _queue_path(self, docx_path: str) -> str:
        base = os.path.basename(docx_path)
        name = os.path.splitext(base)[0]
        return os.path.join(PENDING_DIR, name + ".jsonl")

    # Put item into queue when file is busy
    def _enqueue(self, docx_path: str, payload: str, item_type: str = "text") -> None:
        try:
            os.makedirs(PENDING_DIR, exist_ok=True)
            qpath = self._queue_path(docx_path)
            item = {"ts": time.strftime("%Y-%m-%d %H:%M:%S"), "docx": docx_path, "type": item_type, "payload": payload}
            with open(qpath, "a", encoding="utf-8") as f:
                f.write(json.dumps(item, ensure_ascii=False) + "\n")
            log.info("DocxWriter: item queued in %s", qpath)
        except Exception as e:
            log.exception("DocxWriter: queue error: %s", e)

    # Sync queue for a single .docx
    def sync_queue_for(self, docx_path: str) -> int:
        qpath = self._queue_path(docx_path)
        if not os.path.exists(qpath):
            return 0
        try:
            with open(qpath, "r", encoding="utf-8") as f:
                lines = [ln for ln in f.read().splitlines() if ln.strip()]
        except Exception as e:
            log.exception("DocxWriter: cannot read queue %s: %s", qpath, e)
            return 0
        if not lines:
            os.remove(qpath)
            return 0
        items = []
        for ln in lines:
            try:
                items.append(json.loads(ln))
            except Exception:
                pass
        written = 0
        for item in items:
            payload = item.get("payload", "") or item.get("text", "")
            item_type = item.get("type", "text")
            if not payload:
                continue
            if item_type == "html":
                ok = self.append_html(docx_path, payload)
            elif item_type == "markdown":
                ok = self.append_markdown(docx_path, payload)
            else:
                ok = self.append_text(docx_path, payload)
            if ok:
                written += 1
            else:
                break
        remaining = items[written:]
        try:
            if remaining:
                with open(qpath, "w", encoding="utf-8") as f:
                    for item in remaining:
                        f.write(json.dumps(item, ensure_ascii=False) + "\n")
            else:
                os.remove(qpath)
        except Exception as e:
            log.exception("DocxWriter: queue update error %s: %s", qpath, e)
        if written:
            log.info("DocxWriter: %d items flushed into %s", written, docx_path)
        return written

    # Sync all queues in pending folder
    def sync_all(self) -> int:
        total = 0
        if not os.path.isdir(PENDING_DIR):
            return 0
        for fn in os.listdir(PENDING_DIR):
            if not fn.endswith(".jsonl"):
                continue
            qpath = os.path.join(PENDING_DIR, fn)
            try:
                with open(qpath, "r", encoding="utf-8") as f:
                    first = f.readline().strip()
                if not first:
                    continue
                item = json.loads(first)
                docx_path = item.get("docx", "")
                if docx_path:
                    total += self.sync_queue_for(docx_path)
            except Exception as e:
                log.exception("DocxWriter: sync error %s: %s", qpath, e)
        return total